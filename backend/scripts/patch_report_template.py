"""Bring reportTemplate.xlsx in line with methodology 1.7 and the API.

One-shot, reviewable migration of the Excel template (the workbook itself is
binary, so the formulas live here). Run from the backend directory::

    python scripts/patch_report_template.py            # patch in place
    python scripts/patch_report_template.py --check    # exit 1 if not patched

What changes (methodology section in brackets):

* Inputs: h_reason input; block 2.3 with effective workload values
  (K_calls, m, SP_eff/Prp_eff/A_eff, MRT_hist, SL_pf^raw, R_eff, d_attn^pf,
  d_attn^dec, window check); SL_pf/T_dec/TS formulas of §2.2.
* Reference: generated from the API catalogs (services/report_reference.py).
* Iterations (new sheet): §6.4 steps 1–4, 20 iterations, one state per row,
  convergence status, final state S*, BS=1 state, override state.
* Sizing: d_attn instead of H (§6.1), Th_dec without K_batch, single final
  state for throughput/TTFT/e2e (§6.4–§7), TTFT(1) informational, SLA n/a
  without a target, quotas with K_calls (§9).
* Calibration: η_pf / η_dec without BS, d_attn, MLA-aware M_KV, O_fixed in
  GiB, bottleneck check from two decode runs (Appendix Е).
* Summary / Agentic: labels and checks follow the new single state.
"""

from __future__ import annotations

import argparse
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from services.gpu_catalog_service import load_gpu_catalog  # noqa: E402
from services.llm_catalog_service import load_llm_catalog  # noqa: E402
from services.report_reference import (  # noqa: E402
    GPU_ROWS,
    MODEL_ROWS,
    QUANT_ROWS,
    gpu_display_name,
    write_reference,
)

TEMPLATE = BACKEND / "reportTemplate.xlsx"
MARKER_CELL = ("Inputs", "A80")
MARKER = "2.3. Эффективные значения нагрузки"

OLD_RANGES = {(3, 20): GPU_ROWS, (24, 85): MODEL_ROWS, (89, 94): QUANT_ROWS}

# ── helpers ────────────────────────────────────────────────────────────────


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int = 8) -> None:
    for col in range(1, max_col + 1):
        ws.cell(dst_row, col)._style = copy(ws.cell(src_row, col)._style)


def _set(ws, row: int, label: str, hint: str, kind: str, value, unit: str = "") -> None:
    ws.cell(row, 1).value = label
    ws.cell(row, 2).value = hint
    ws.cell(row, 3).value = kind
    ws.cell(row, 4).value = value
    ws.cell(row, 5).value = unit


def _sel(a: str, b: str) -> str:
    """select_th_*: min of two branches, or the only positive one."""
    return f"IF(AND({a}<=0,{b}<=0),0,IF({b}<=0,{a},IF({a}<=0,{b},MIN({a},{b}))))"


def _remap_reference_ranges(wb) -> None:
    pattern = re.compile(r"Reference!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)")

    def repl(m: re.Match) -> str:
        first, last = int(m.group(2)), int(m.group(4))
        new = OLD_RANGES.get((first, last))
        if new is None:
            return m.group(0)
        return f"Reference!${m.group(1)}${new[0]}:${m.group(3)}${new[1]}"

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = pattern.sub(repl, cell.value)
        for dv in ws.data_validations.dataValidation:
            if dv.formula1:
                dv.formula1 = pattern.sub(repl, dv.formula1)


# ── Reference ──────────────────────────────────────────────────────────────


def patch_reference(wb) -> None:
    ws = wb["Reference"]
    styles = {
        "title": copy(ws["A1"]._style),
        "header": copy(ws["A2"]._style),
        "data": copy(ws["B3"]._style),
    }
    label_style = copy(ws["A3"]._style)
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in range(1, max(ws.max_row, QUANT_ROWS[1]) + 1):
        for col in range(1, 17):
            ws.cell(row, col).value = None
    write_reference(ws, load_gpu_catalog(), load_llm_catalog(), styles)
    for first, last in (GPU_ROWS, MODEL_ROWS, QUANT_ROWS):
        for row in range(first, last + 1):
            ws.cell(row, 1)._style = copy(label_style)
    _remap_reference_ranges(wb)


# ── Inputs ─────────────────────────────────────────────────────────────────


def patch_inputs(wb) -> None:
    ws = wb["Inputs"]
    # Default selections must exist in the generated Reference lists.
    h100 = next(g for g in load_gpu_catalog() if g["id"] == "nvidia-h100-gpu-accelerator-sxm-card")
    ws["D25"] = gpu_display_name(h100)

    ws["A20"] = "Полная длина сессии (TS)"
    ws["B20"] = "TS = SL_pf^raw + T_dec (§2.2). Для KV-кэша берётся SL = min(TS, TS_max)."
    ws["D20"] = "=D88+D22"
    ws["A21"] = "Длина входа на этапе prefill (SL_pf)"
    ws["B21"] = "SL_pf = min(SL_pf^raw, TS_max − T_dec) (§2.2). Для TTFT, Th_pf и Th_dec."
    ws["D21"] = "=MIN(D88,D44-D22)"
    ws["A22"] = "Число токенов decode (T_dec)"
    ws["B22"] = "T_dec = A_eff + MRT (§2.2, В.3): A_eff = A + A_tool."
    ws["D22"] = "=D85+D86"

    _copy_row_style(ws, 19, 23)
    _set(
        ws,
        23,
        "Reasoning прошлых ходов в истории (h_reason)",
        "0 — движок отбрасывает reasoning прошлых ходов (по умолчанию); 1 — сохраняет (§2.2)",
        "INPUT",
        0,
        "0/1",
    )
    dv = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
    dv.add("D23")
    ws.add_data_validation(dv)

    # Block 2.3: effective workload values (§2.2, Appendix В.3).
    ws.merge_cells("A80:H80")
    ws["A80"] = (
        "2.3. Эффективные значения нагрузки (§2.2, Приложение В.3) — рассчитываются "
        "автоматически из Inputs и Agentic"
    )
    ws["A80"]._style = copy(ws["A74"]._style)
    rows = [
        (
            "Вызовов LLM на ход (K_calls)",
            "Agentic!D23; 1 для неагентного сценария",
            "=MAX(1,Agentic!D23)",
            "pcs",
        ),
        ("Число LLM-вызовов в истории (m)", "m = N_prp · K_calls", "=D19*D81", "pcs"),
        (
            "Системный промпт эффективный (SP_eff)",
            "SP + SP_tools + C_rag_static",
            "=Agentic!D30",
            "tokens",
        ),
        ("Запрос эффективный (Prp_eff)", "Prp + C_rag_dynamic", "=Agentic!D31", "tokens"),
        ("Ответ эффективный (A_eff)", "A + A_tool", "=Agentic!D32", "tokens"),
        (
            "Reasoning текущего вызова (MRT)",
            "Генерируется на каждом вызове",
            "=Agentic!D33",
            "tokens",
        ),
        ("Reasoning в истории (h_reason · MRT)", "0 при h_reason = 0", "=D23*D86", "tokens"),
        (
            "Вход prefill до окна (SL_pf^raw)",
            "SP_eff + m·Prp_eff + (m−1)·(A_eff + h_reason·MRT)",
            "=D83+D82*D84+MAX(0,D82-1)*(D85+D87)",
            "tokens",
        ),
        (
            "Интенсивность вызовов (R_eff = R · K_calls)",
            "Используется в §6.4 и §9",
            "=D75*D81",
            "req/s",
        ),
        (
            "Размерность attention prefill (d_attn^pf)",
            "MHA/GQA/MQA: N_attn·head_dim; MLA: N_attn·(2·qk_head_dim − rope)/2 (§6.1)",
            "=IF(D47,D40*(2*D42-D46)/2,D40*IF(D42>0,D42,D39/D40))",
            "-",
        ),
        (
            "Размерность attention decode (d_attn^dec)",
            "MHA/GQA/MQA: = d_attn^pf; MLA: N_attn·(2·kv_lora_rank + rope)/2 (§6.1)",
            "=IF(D47,D40*(2*D45+D46)/2,D90)",
            "-",
        ),
        (
            "Проверка контекстного окна",
            "T_dec < TS_max обязательно; при SL_pf < SL_pf^raw история усечена окном",
            '=IF(D22>=D44,"✗ FAIL: T_dec ≥ TS_max — генерация не помещается в окно",'
            'IF(D21<D88,"⚠ история усечена окном: SL_pf < SL_pf^raw","✓ OK"))',
            "status",
        ),
    ]
    for offset, (label, hint, formula, unit) in enumerate(rows):
        row = 81 + offset
        _copy_row_style(ws, 20, row)
        _set(ws, row, label, hint, "calc", formula, unit)


# ── Agentic ────────────────────────────────────────────────────────────────


def patch_agentic(wb) -> None:
    ws = wb["Agentic"]
    ws["B34"] = "TS = SL_pf^raw + T_dec (Inputs!D20, §2.2)"
    ws["D34"] = "=Inputs!D20"
    ws["B35"] = "SL_pf после окна (Inputs!D21)"
    ws["D35"] = "=Inputs!D21"
    for row, ref in ((39, "D15+SP_tools"), (40, "D16+C_rag"), (41, "D17+A_tool"), (42, "MRT")):
        ws[f"B{row}"] = f"Учитывается автоматически в Inputs (блок 2.3): {ref}"
    ws["B43"] = "R_eff = R · K_calls — учитывается автоматически (Inputs!D89)"
    ws["B46"] = "SL_pf после окна, учитывается автоматически"
    ws["D46"] = "=Inputs!D21"


# ── Iterations (new sheet) ─────────────────────────────────────────────────

FIRST_ITER_ROW = 6
N_ITER = 20
LAST_ITER_ROW = FIRST_ITER_ROW + N_ITER - 1  # k = 0..19
STATUS_ROW = 28
CAND_ROW = 30  # candidate state (oscillation / non-convergence)
FINAL_ROW = 33  # S*
BS1_ROW = 34  # BS = 1 (unloaded)
OVERRIDE_ROW = 35  # Servers_final (override aware)

HEADERS = [
    "k",
    "Servers^(k)",
    "BS_real",
    "P_eff(BS)",
    "P_eff(BS+1)",
    "Th_pf^cb,mem",
    "Th_pf",
    "Th_dec^comp/сесс",
    "Th_dec^mem/сесс",
    "Th_dec",
    "C^model",
    "Th_server",
    "Servers^comp",
    "Servers^(k+1)",
    "Сошлось",
]


def _state_formulas(r: int, servers: str, bs_override: str | None = None) -> dict[str, str]:
    i = "Inputs!"
    bs = (
        bs_override
        or f"=IF(AND(Sizing!$D$12>0,B{r}>0),"
        f"MIN(IF(Sizing!$D$15>0,Sizing!$D$15,10^9),"
        f"MAX(1,CEILING({i}$D$12/(Sizing!$D$12*B{r}),1))),1)"
    )

    def p_eff(bs_expr: str) -> str:
        return f"=IF({i}$D$52,{i}$D$48+{i}$D$49*(1-POWER(1-{i}$D$50/{i}$D$51,{bs_expr})),{i}$D$37)"

    dec_cmp = (
        f"=IF(C{r}>0,Sizing!$D$26*{i}$D$62/(Sizing!$D$25+4*{i}$D$38*{i}$D$91*"
        f"({i}$D$21+({i}$D$22-1)/2))/C{r},0)"
    )
    mem_den = f"(%s*10^9*{i}$D$54+C{r}*Sizing!$D$7*1024^3+{i}$D$72*1024^3)"
    return {
        "B": servers if servers.startswith("=") else f"={servers}",
        "C": bs,
        "D": p_eff(f"C{r}"),
        "E": p_eff(f"C{r}+1"),
        "F": (
            f"=IF(AND({i}$D$28>0,{i}$D$70>0),{i}$D$70*{i}$D$28*10^9*{i}$D$63/"
            + (mem_den % f"E{r}")
            + ",0)"
        ),
        "G": (
            f'=IF({i}$D$65>0,{i}$D$65,IF({i}$D$69="continuous",'
            + _sel("Sizing!$D$30", f"F{r}")
            + ",Sizing!$D$29))"
        ),
        "H": dec_cmp,
        "I": f"=IF({i}$D$28>0,{i}$D$28*10^9*{i}$D$63/" + (mem_den % f"D{r}") + ",0)",
        "J": f"=IF({i}$D$66>0,{i}$D$66,{_sel(f'H{r}', f'I{r}')})*{i}$D$68",
        "K": f"=IF(AND(G{r}>0,J{r}>0),C{r}/(Sizing!$D$27/G{r}+{i}$D$22/J{r}),0)",
        "L": f"=Sizing!$D$12*K{r}",
        "M": f"=IF(L{r}>0,CEILING({i}$D$12*{i}$D$89*{i}$D$76/L{r},1),0)",
        "N": f"=MAX(Sizing!$D$22,M{r})",
        "O": f"=N{r}=B{r}",
    }


def patch_iterations(wb) -> None:
    if "Iterations" in wb.sheetnames:
        del wb["Iterations"]
    sizing = wb["Sizing"]
    ws = wb.create_sheet("Iterations", index=wb.sheetnames.index("Sizing") + 1)
    title_style = copy(sizing["A1"]._style)
    section_style = copy(sizing["A24"]._style)
    header_style = copy(sizing["A2"]._style)
    calc_style = copy(sizing["D25"]._style)
    label_style = copy(sizing["A25"]._style)

    ws["A1"] = "§6.4 Итеративная резолюция BS_real ↔ Servers (шаги 1–4, k_max = 20)"
    ws["A1"]._style = title_style
    ws["A2"] = (
        "Каждая строка — одно состояние: Servers^(k) → BS_real → Th_pf(BS), Th_dec(BS) → "
        "C^model(BS) → Servers^comp → Servers^(k+1) = max(Servers_mem, Servers^comp). "
        "Сходимость — первая строка, где Servers^(k+1) = Servers^(k) (шаг 3 §6.4)."
    )
    ws["A4"] = "Итерации"
    ws["A4"]._style = section_style
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(5, col)
        cell.value = title
        cell._style = header_style

    for k in range(N_ITER):
        r = FIRST_ITER_ROW + k
        servers = "=Sizing!$D$22" if k == 0 else f"=N{r - 1}"
        ws.cell(r, 1).value = k
        ws.cell(r, 1)._style = label_style
        for col, formula in _state_formulas(r, servers).items():
            cell = ws[f"{col}{r}"]
            cell.value = formula
            cell._style = calc_style

    first, last = FIRST_ITER_ROW, LAST_ITER_ROW
    ws[f"A{STATUS_ROW}"] = "Статус сходимости"
    ws[f"A{STATUS_ROW}"]._style = label_style
    ws[f"B{STATUS_ROW}"] = (
        f'=IF(ISNUMBER(MATCH(TRUE(),O{first}:O{last},0)),"converged",'
        f'IF(N{CAND_ROW}<=B{CAND_ROW},"oscillating","not_converged"))'
    )
    ws[f"A{STATUS_ROW + 1}"] = "Число итераций до сходимости"
    ws[f"A{STATUS_ROW + 1}"]._style = label_style
    ws[f"B{STATUS_ROW + 1}"] = f"=IFERROR(MATCH(TRUE(),O{first}:O{last},0),{N_ITER})"

    ws[f"A{CAND_ROW - 1}"] = (
        "Кандидат без сходимости: max(Servers^(20), Servers^(19)); самосогласован, если "
        "max(Servers_mem, Servers^comp(кандидат)) ≤ кандидат"
    )
    ws.cell(CAND_ROW, 1).value = "кандидат"
    ws.cell(CAND_ROW, 1)._style = label_style
    for col, formula in _state_formulas(CAND_ROW, f"=MAX(N{last},B{last})").items():
        ws[f"{col}{CAND_ROW}"] = formula
        ws[f"{col}{CAND_ROW}"]._style = calc_style

    ws[f"A{FINAL_ROW - 1}"] = "Итоговые состояния"
    ws[f"A{FINAL_ROW - 1}"]._style = section_style
    s_star = (
        f'=IF(B{STATUS_ROW}="converged",INDEX(B{first}:B{last},B{STATUS_ROW + 1}),'
        f'IF(B{STATUS_ROW}="oscillating",B{CAND_ROW},MAX(B{CAND_ROW},N{CAND_ROW})))'
    )
    states = (
        (FINAL_ROW, "S* (сошедшееся)", s_star, None),
        (BS1_ROW, "BS = 1 (ненагруженный)", f"=B{FINAL_ROW}", "=1"),
        (OVERRIDE_ROW, "Servers_final (с override)", "=Sizing!$D$68", None),
    )
    for row, label, servers, bs in states:
        ws.cell(row, 1).value = label
        ws.cell(row, 1)._style = label_style
        for col, formula in _state_formulas(row, servers, bs).items():
            ws[f"{col}{row}"] = formula
            ws[f"{col}{row}"]._style = calc_style
    ws.column_dimensions["A"].width = 28
    for col in "BCDEFGHIJKLMNO":
        ws.column_dimensions[col].width = 15


# ── Sizing ─────────────────────────────────────────────────────────────────


def patch_sizing(wb) -> None:
    ws = wb["Sizing"]
    it = "Iterations!"
    i = "Inputs!"
    fr, b1, ov = FINAL_ROW, BS1_ROW, OVERRIDE_ROW
    pf_den = f"(D25+4*{i}D38*{i}D90*{i}D21)"

    ws["A6"] = "SL для KV-кэша (min(TS, TS_max))"
    ws["B6"] = "SL = SL_pf + T_dec = min(TS, TS_max) (§2.2)"
    ws["A27"] = "SL_pf эффективная (после префикс-кэша)"
    ws["A28"] = "BS_real итоговый (при S*)"
    ws["B28"] = "BS_real на сошедшемся состоянии §6.4 (лист Iterations)"
    ws["D28"] = f"={it}C{fr}"
    ws["B29"] = "F_count · η_pf · K_batch / (FPS + 4·L·d_attn^pf·SL_pf). Только для static."
    ws["D29"] = f"=D26*{i}D61*D18/{pf_den}"
    ws["B30"] = "F_count · η_pf / (FPS + 4·L·d_attn^pf·SL_pf); attention по полному SL_pf (§6.1)"
    ws["D30"] = f"=D26*{i}D61/{pf_den}"
    ws["B31"] = "C_pf·BW·10⁹·η_mem / (P_eff(BS+1)·10⁹·B_quant + BS·M_KV·1024³ + O_fixed·1024³)"
    ws["D31"] = f"={it}F{fr}"
    ws["B32"] = "min(compute, mem) при BS_real итоговом"
    ws["D32"] = f"={_sel('D30', 'D31')}"
    ws["D33"] = f"={it}G{fr}"
    ws["D34"] = f"=IF(D26>0,D33*{pf_den}/(D26*{i}D61),0)"
    ws["B35"] = (
        "F_count · η_dec / (FPS + 4·L·d_attn^dec·(SL_pf + (T_dec−1)/2)) / BS_real; без K_batch"
    )
    ws["D35"] = f"={it}H{fr}"
    ws["D36"] = f"={it}I{fr}"
    ws["D37"] = f"={_sel('D35', 'D36')}"
    ws["D38"] = f"={it}J{fr}"
    ws["B39"] = "BS_real / (SL_pf^eff/Th_pf(BS) + T_dec/Th_dec(BS)) при BS_real итоговом"
    ws["D39"] = f"={it}K{fr}"
    ws["D40"] = f"={it}L{fr}"
    ws["B41"] = "⌈S_sim · R_eff · K_SLA / Th_server_comp⌉ при S*; R_eff = R · K_calls"
    ws["D41"] = f"={it}M{fr}"

    ws["A42"] = "§6.4 Итеративная резолюция BS_real ↔ Servers (полная таблица — лист Iterations)"
    for k in range(10):
        r = 43 + k
        ws[f"A{r}"] = f"Servers^({k})"
        ws[f"B{r}"] = f"Iterations: состояние k = {k}"
        ws[f"D{r}"] = f"={it}B{FIRST_ITER_ROW + k}"
    ws["A53"] = "Servers_comp (на S*; при отсутствии сходимости = S*)"
    ws["B53"] = "Servers^comp(S*) при converged, иначе S* (§6.4 шаг 4)"
    ws["D53"] = f'=IF({it}B{STATUS_ROW}="converged",{it}M{fr},{it}B{fr})'
    ws["A54"] = "Статус сходимости §6.4"
    ws["B54"] = "converged / oscillating / not_converged — шаги 3–4 §6.4"
    ws["D54"] = (
        f'=IF({it}B{STATUS_ROW}="converged","✓ OK: converged за "&{it}B{STATUS_ROW + 1}&" итер.",'
        f'IF({it}B{STATUS_ROW}="oscillating","⚠ oscillating: взят больший из двух последних",'
        f'"⚠ not_converged: проверьте Servers_override"))'
    )
    ws["A55"] = "BS_real (сошедшийся)"
    ws["D55"] = f"={it}C{fr}"
    ws["B56"] = (
        "N_TP=Z · BW·10⁹·η_mem / (T_dec·M_KV·1024³) — справочно; применим только в "
        "decode-dominated режиме (§6.4)"
    )

    ws["A57"] = "ρ_pf (загрузка prefill, диагностика §6.2)"
    ws["B57"] = (
        "C^model · SL_pf^eff / Th_pf — среднее число запросов в prefill; > 1 → prefill в "
        "очереди, C^model завышен (проверить нагрузочным тестом / PD)"
    )
    ws["C57"] = "calc"
    ws["D57"] = "=IF(D33>0,D39*D27/D33,0)"
    ws["E57"] = "-"
    for col in "ABCDE":
        ws[f"{col}57"]._style = copy(ws[f"{col}56"]._style)

    ws["A59"] = "TTFT при BS_real итоговом"
    ws["B59"] = "SL_pf^eff / Th_pf(BS) + 1/Th_dec(BS) + T_overhead (§7.1)"
    ws["D59"] = f"=IF(AND(D33>0,D38>0),D27/D33+1/D38+{i}D71,0)"
    ws["A60"] = "GenerationTime при BS_real итоговом"
    ws["D60"] = "=IF(D38>0,Inputs!D22/D38,0)"
    ws["A61"] = "e2eLatency_analyt (BS = 1, ненагруженная)"
    ws["B61"] = "TTFT(1) + T_dec/Th_dec(1) — справочно (§7.2)"
    ws["D61"] = f"=IF({it}J{b1}>0,D65+{i}D22/{it}J{b1},0)"
    ws["B62"] = "max(e2e_analyt, BS_real / C^model(BS)) — используется для SLA (§7.3)"
    ws["D62"] = "=IF(D39>0,MAX(D61,D28/D39),D61)"
    ws["B63"] = "TTFT_SLA ≥ TTFT(BS итоговый); n/a без цели"
    ws["D63"] = f'=IF({i}D77>0,IF({i}D77>=D59,"✓ OK","✗ FAIL"),"n/a")'
    ws["B64"] = "e2eLatency_SLA ≥ e2eLatency_load; n/a без цели"
    ws["D64"] = f'=IF({i}D78>0,IF({i}D78>=D62,"✓ OK","✗ FAIL"),"n/a")'
    for col in "ABCDE":
        ws[f"{col}65"]._style = copy(ws[f"{col}60"]._style)
    ws["A65"] = "TTFT(BS = 1) — справочно"
    ws["B65"] = "SL_pf^eff / Th_pf(1) + 1/Th_dec(1) + T_overhead"
    ws["C65"] = "calc"
    ws["D65"] = f"=IF(AND({it}G{b1}>0,{it}J{b1}>0),D27/{it}G{b1}+1/{it}J{b1}+{i}D71,0)"
    ws["E65"] = "s"

    ws["A67"] = "Серверов по расчёту (S*, §6.4)"
    ws["B67"] = "max(Servers_mem, Servers_comp) на сошедшемся состоянии"
    ws["D67"] = f"={it}B{fr}"
    ws["D70"] = f"={it}C{ov}"
    ws["A71"] = "Th_dec при Servers_final"
    ws["B71"] = "min(compute, mem) · K_spec при BS_real(Servers_final)"
    ws["D71"] = f"={it}J{ov}"
    ws["A72"] = "e2eLatency_load при Servers_final"
    ws["B72"] = "max(e2e_analyt, BS_real / C^model) при Servers_final"
    ws["D72"] = f"=IF({it}K{ov}>0,MAX(D61,D70/{it}K{ov}),D61)"
    fail_hint = (
        "✗ FAIL — см. Приложение Б: Б.1 увеличить Z, Б.2 квантизация (FP8/INT4), "
        "Б.3 сократить SL, Б.4 сократить T_dec, Б.5 меньшая модель, Б.6 лучший GPU, "
        "Б.7 смягчить SLA"
    )
    ttft_ov = f"IF(AND({it}G{ov}>0,{it}J{ov}>0),D27/{it}G{ov}+1/{it}J{ov}+{i}D71,0)"
    ws["B73"] = "TTFT и e2e при Servers_final против целей; n/a без целей"
    sla_ok = f"AND(OR({i}D77<=0,{i}D77>={ttft_ov}),OR({i}D78<=0,{i}D78>=D72))"
    ws["D73"] = (
        f'=IF(D13<>"✓ OK","✗ FAIL: модель не помещается на сервер",'
        f'IF(LEFT({i}D92,1)="✗",{i}D92,'
        f'IF(AND({i}D77<=0,{i}D78<=0),"n/a: цели SLA не заданы",'
        f'IF({sla_ok},"✓ OK: конфигурация удовлетворяет SLA","{fail_hint}"))))'
    )

    ws["B76"] = "S_sim · R_eff · K_SLA · 60; R_eff = R · K_calls (§9)"
    ws["D76"] = f"={i}D12*{i}D89*{i}D76*60"
    ws["B77"] = "S_sim · R_eff · 60"
    ws["D77"] = f"={i}D12*{i}D89*60"


# ── Calibration ────────────────────────────────────────────────────────────


def patch_calibration(wb) -> None:
    ws = wb["Calibration"]
    i = "Inputs!"
    ws["B11"] = "M_KV под SL прогона = prompt + output; MLA: L·(r + rope)·SL·B_state"
    ws["D11"] = (
        f"=IF({i}D47,{i}D38*({i}D45+{i}D46)*(D9+D10)*{i}D55*{i}D57/1024^3,"
        f"2*{i}D38*{i}D41*{i}D42*(D9+D10)*{i}D55*{i}D57/1024^3)"
    )
    ws["B12"] = (
        "Е.2: η_mem = Th₁/(N·BS) · (P_eff(BS)·10⁹·B_q + BS·M_KV·1024³ + O_fixed·1024³) / (BW·10⁹)"
    )
    ws["D12"] = (
        f"=IF(AND(D7>0,{i}D28>0,Sizing!D12>0),D8*(IF({i}D52,({i}D48+{i}D49*(1-POWER(1-"
        f"{i}D50/{i}D51,D7))),{i}D37)*10^9*{i}D54+D7*D11*1024^3+{i}D72*1024^3)/"
        f"({i}D28*10^9*D7*Sizing!D12),0)"
    )
    ws["B13"] = (
        "Е.3: η_dec = Th_dec^obs · (FPS + 4·L·d_attn^dec·(SL + (T−1)/2)) / (F_count · N). "
        "Без множителя BS; валиден только при compute-bound decode."
    )
    ws["D13"] = (
        f"=IF(AND(Sizing!D26>0,Sizing!D12>0),D8*(Sizing!D25+4*{i}D38*{i}D91*(D9+(D10-1)/2))"
        f"/(Sizing!D26*Sizing!D12),0)"
    )
    ws["A14"] = "Узкое место decode (по двум прогонам)"
    ws["B14"] = (
        "Е.3: по одному прогону bottleneck не определяется. Сравните прогоны 1 и 2: "
        "aggregate Th_dec растёт ~∝ BS → memory-bound (калибруйте η_mem); почти не растёт → "
        "compute-bound (η_dec валиден)."
    )
    ws["D14"] = (
        "=IF(AND(D7>0,D18>0,D8>0,D19>0,D7<>D18),"
        "IF(LN(MAX(D8,D19)/MIN(D8,D19))/LN(MAX(D7,D18)/MIN(D7,D18))>=0.5,"
        '"memory-bound: используйте η_mem; η_dec не калибруется",'
        '"compute-bound: η_dec валиден"),'
        '"нужен второй прогон с другим BS")'
    )
    ws["A24"] = "O_fixed (выведено из 2 прогонов)"
    ws["B24"] = "O = (t₂·D₂ − t₁·D₁) / (1024³ · (t₁ − t₂)); в GiB. Если ≤ 0 → O_fixed = 0."
    ws["E24"] = "GiB"
    ws["B31"] = (
        "Е.3: η_pf = Th_pf^obs · (FPS + 4·L·d_attn^pf·SL_pf) / (F_count · N). Без делителя BS."
    )
    ws["D31"] = (
        f"=IF(AND(Sizing!D26>0,Sizing!D12>0),D29*(Sizing!D25+4*{i}D38*{i}D90*D30)"
        f"/(Sizing!D26*Sizing!D12),0)"
    )


# ── Summary ────────────────────────────────────────────────────────────────


def patch_summary(wb) -> None:
    ws = wb["Summary"]
    ws["B11"] = (
        '=IF(Agentic!D23<=1,"✓ n/a: K_calls=1 (не агентный сценарий)",'
        '"✓ OK: агентные надбавки и K_calls учтены автоматически (Inputs, блок 2.3)")'
    )
    ws["B12"] = '=IF(ISNUMBER(SEARCH("OK",Sizing!D54)),"✓ OK: итерации сошлись",Sizing!D54)'
    ws["B49"] = (
        '=IF(Inputs!D77<=0,"n/a: цель TTFT не задана",'
        'IF(Inputs!D77>=Sizing!D59,"✓ OK: в пределах SLA","✗ FAIL: превышение SLA"))'
    )
    ws["B50"] = (
        '=IF(Inputs!D78<=0,"n/a: цель e2e не задана",'
        'IF(Inputs!D78>=Sizing!D62,"✓ OK: в пределах SLA","✗ FAIL: превышение SLA"))'
    )
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if isinstance(label, str) and "(10)" in label:
            ws.cell(row, 1).value = label.replace("(10)", "(S*)")


def is_patched(wb) -> bool:
    sheet, cell = MARKER_CELL
    value = wb[sheet][cell].value
    return isinstance(value, str) and value.startswith(MARKER)


def patch(wb) -> None:
    patch_reference(wb)
    patch_inputs(wb)
    patch_agentic(wb)
    patch_iterations(wb)
    patch_sizing(wb)
    patch_calibration(wb)
    patch_summary(wb)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--template", type=Path, default=TEMPLATE)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.template)
    if args.check:
        return 0 if is_patched(wb) else 1
    if is_patched(wb):
        print(f"{args.template} already patched")
        return 0
    patch(wb)
    out = args.output or args.template
    wb.save(out)
    print(f"patched -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
