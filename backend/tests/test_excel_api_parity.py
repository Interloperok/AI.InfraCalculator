"""Excel report recomputed by LibreOffice must match the API (methodology 1.7).

Generates the report for a scenario matrix, recalculates every workbook with a
headless LibreOffice and compares the key cells with ``run_sizing``. Skipped
when ``soffice`` is not installed (CI images without LibreOffice).
"""

from __future__ import annotations

import copy
import os
import shutil
import subprocess
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from models import SizingInput  # noqa: E402
from services.gpu_catalog_service import lookup_gpu_tflops  # noqa: E402
from services.report_service import ReportGenerator  # noqa: E402
from services.sizing_service import run_sizing  # noqa: E402

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
pytestmark = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice (soffice) not installed")

Q32 = dict(
    model_name="Qwen3-32B", params_billions=32, params_active=32, layers_L=64,
    hidden_size_H=5120, num_attention_heads=64, num_kv_heads=8, head_dim=128,
    max_context_window_TSmax=40960,
)  # fmt: skip
H200 = dict(
    gpu_id="nvidia-h200-gpu-accelerator-sxm-card", gpu_name="H200 SXM 141GB", gpu_mem_gb=141,
    gpus_per_server=4, bw_gpu_gbs=4800, gpu_flops_Fcount=989,
)  # fmt: skip
BASE = dict(
    internal_users=1000, penetration_internal=0.6, concurrency_internal=0.15, external_users=0,
    penetration_external=0, concurrency_external=0, sessions_per_user_J=1,
    system_prompt_tokens_SP=1000, user_prompt_tokens_Prp=1400, answer_tokens_A=400,
    reasoning_tokens_MRT=0, dialog_turns=3, bytes_per_param=1, bytes_per_kv_state=1,
    safe_margin=5, emp_model=1, emp_kv=1, kavail=0.9, tp_multiplier_Z=1, saturation_coeff_C=6,
    eta_prefill=0.167, eta_decode=0.2, eta_mem=0.36, eta_cache=0.5, k_spec=1.8,
    engine_mode="continuous", c_pf=256, t_overhead=0.026, o_fixed=0, rps_per_session_R=0.02,
    sla_reserve_KSLA=1.25, ttft_sla=3, e2e_latency_sla=30, **Q32, **H200,
)  # fmt: skip
QMOE = dict(
    model_name="Qwen3-30B-A3B", params_billions=30, params_active=3, layers_L=48,
    hidden_size_H=2048, num_attention_heads=32, num_kv_heads=4, head_dim=128,
    max_context_window_TSmax=40960, params_dense=1.54, params_moe=28.99, k_experts=8,
    n_experts=128,
)  # fmt: skip
DSV3 = dict(
    model_name="DeepSeek-V3", params_billions=671, params_active=37, layers_L=61,
    hidden_size_H=7168, num_attention_heads=128, num_kv_heads=128, head_dim=192,
    max_context_window_TSmax=163840, params_dense=17.12, params_moe=653.91, k_experts=8,
    n_experts=256, kv_lora_rank=512, qk_rope_head_dim=64, gpus_per_server=8,
)  # fmt: skip


def _v(**kw):
    d = copy.deepcopy(BASE)
    d.update(kw)
    return d


SCENARIOS = {
    "base": _v(),
    "agentic_k3": _v(
        user_prompt_tokens_Prp=200,
        c_rag_dynamic=1200,
        k_calls=3,
        sp_tools=500,
        c_rag_static=300,
        a_tool=100,
    ),  # fmt: skip
    "static": _v(engine_mode="static"),
    "moe": _v(**QMOE),
    "mla": _v(**DSV3),
    "reasoning_in_history": _v(reasoning_tokens_MRT=1500, reasoning_in_history=True),
    "window_truncated": _v(dialog_turns=30, bytes_per_param=2, max_context_window_TSmax=16384),
    "heavy_moe_reasoning": _v(**QMOE, internal_users=20000, reasoning_tokens_MRT=2000, eta_cache=0),
    "empirical": _v(th_prefill_empir=20000, th_decode_empir=40),
    "no_sla": _v(ttft_sla=None, e2e_latency_sla=None),
}

CELLS = {
    "servers_by_memory": "Sizing!D22",
    "servers_by_compute": "Sizing!D53",
    "servers_final": "Sizing!D68",
    "BS_real": "Sizing!D28",
    "iteration_count": "Iterations!B29",
    "iteration_status": "Iterations!B28",
    "th_prefill": "Sizing!D33",
    "th_decode": "Sizing!D38",
    "Cmodel_rps": "Sizing!D39",
    "ttft_analyt": "Sizing!D59",
    "ttft_bs1": "Sizing!D65",
    "e2e_latency_analyt": "Sizing!D61",
    "e2e_latency_load": "Sizing!D62",
    "peak_rpm": "Sizing!D76",
    "peak_tpm_input": "Sizing!D78",
    "peak_tpm_output": "Sizing!D79",
    "SL_sequence_length": "Sizing!D6",
    "SL_pf_input_length": "Inputs!D21",
    "kv_per_session_gb": "Sizing!D7",
    "rho_pf": "Sizing!D57",
}
REL_TOL = 1e-3  # API rounds published floats to 4 decimals


@pytest.fixture(scope="module")
def recalculated(tmp_path_factory):
    work = tmp_path_factory.mktemp("parity")
    gen = ReportGenerator(gpu_tflops_lookup=lookup_gpu_tflops)
    paths = []
    for name, payload in SCENARIOS.items():
        path = work / f"{name}.xlsx"
        path.write_bytes(gen.generate(SizingInput(**payload)).getvalue())
        paths.append(str(path))
    out = work / "calc"
    out.mkdir()
    subprocess.run(
        [
            SOFFICE, "--headless", f"-env:UserInstallation=file://{work}/lo", "--calc",
            "--convert-to", "xlsx", "--outdir", str(out), *paths,
        ],
        check=True, capture_output=True, timeout=600,
    )  # fmt: skip
    return {
        name: openpyxl.load_workbook(out / f"{name}.xlsx", data_only=True) for name in SCENARIOS
    }


def _cell(wb, ref):
    sheet, cell = ref.split("!")
    return wb[sheet][cell].value


@pytest.mark.parametrize("name", list(SCENARIOS))
def test_excel_matches_api(recalculated, name):
    # Формулы Excel — §6.4 без подбора под SLA, проверка SLA одного вызова
    api = run_sizing(SizingInput(**{**SCENARIOS[name], "sla_fit_servers": False, "e2e_sla_scope": "call"})).model_dump()
    wb = recalculated[name]
    final = run_sizing(SizingInput(**SCENARIOS[name]))
    ws = wb["Итог API"]
    assert ws["B5"].value == final.servers_final and ws["B15"].value == final.sizing_status
    for key, ref in CELLS.items():
        expected, actual = api.get(key), _cell(wb, ref)
        if expected is None:
            continue
        if isinstance(expected, str):
            assert actual == expected, f"{key}: Excel {actual!r} != API {expected!r}"
            continue
        assert isinstance(actual, (int, float)), f"{key}: Excel {actual!r}"
        assert actual == pytest.approx(expected, rel=REL_TOL, abs=1e-4), key
    verdict = str(_cell(wb, "Sizing!D73"))
    if api["ttft_sla_target"] is None and api["e2e_latency_sla_target"] is None:
        assert verdict.startswith("n/a")
    else:
        assert verdict.startswith("✓") == bool(api["sla_passed"]), verdict
