"""Committed Excel template matches methodology 1.7 layout and API catalogs."""

from __future__ import annotations

import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from scripts.patch_report_template import TEMPLATE, is_patched  # noqa: E402
from services.gpu_catalog_service import load_gpu_catalog  # noqa: E402
from services.llm_catalog_service import load_llm_catalog  # noqa: E402
from services.report_reference import (  # noqa: E402
    GPU_ROWS,
    MODEL_ROWS,
    QUANT_ROWS,
    datacenter_gpus,
    gpu_row,
    model_row,
)
from core.methodology_constants import QUANTIZATION_FORMATS  # noqa: E402


def _wb():
    return openpyxl.load_workbook(TEMPLATE)


def _block(ws, rows, width):
    first, last = rows
    out = []
    for r in range(first, last + 1):
        if ws.cell(r, 1).value is None:
            break
        out.append([ws.cell(r, c).value for c in range(1, width + 1)])
    return out


def test_template_is_patched():
    assert is_patched(_wb())


def test_reference_matches_catalogs():
    ref = _wb()["Reference"]
    assert _block(ref, GPU_ROWS, 6) == [gpu_row(g) for g in datacenter_gpus(load_gpu_catalog())]
    assert _block(ref, MODEL_ROWS, 16) == [model_row(m) for m in load_llm_catalog()]
    assert [r[0] for r in _block(ref, QUANT_ROWS, 3)] == [q["label"] for q in QUANTIZATION_FORMATS]


def test_formulas_reference_existing_sheets_only():
    wb = _wb()
    sheets = set(wb.sheetnames)
    ref_re = re.compile(r"([A-Za-z]+)!\$?[A-Z]+\$?\d+")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    assert "#REF" not in value, f"{ws.title}!{cell.coordinate}"
                    for sheet in ref_re.findall(value):
                        assert sheet in sheets, f"{ws.title}!{cell.coordinate} -> {sheet}"


def test_no_k_batch_and_no_hidden_size_in_throughput_formulas():
    sizing = _wb()["Sizing"]
    # Th_pf continuous / roofline use d_attn^pf (Inputs!D90), not H (Inputs!D39).
    for cell in ("D30", "D34"):
        assert "Inputs!D90" in sizing[cell].value and "Inputs!D39" not in sizing[cell].value
    # Iterations: Th_dec compute branch has no K_batch (Sizing!D18) and uses d_attn^dec.
    it = _wb()["Iterations"]
    assert "D$91" in it["H6"].value and "D18" not in it["H6"].value


def test_quotas_scale_with_k_calls():
    sizing = _wb()["Sizing"]
    assert "Inputs!D89" in sizing["D76"].value
