"""Reference sheet of the Excel report, generated from the API catalogs.

Single source of truth: the GPU list comes from ``gpu_data.json`` + the
``gpu_overrides.json`` overlay, the model list from ``llm_data.json`` and the
quantization list from ``core.methodology_constants.QUANTIZATION_FORMATS``.
The same function fills the committed template (scripts/patch_report_template.py)
and refreshes the lists on every download, so the workbook never carries
hand-typed catalog numbers that drift from the API.

Layout (rows are fixed so the Inputs INDEX/MATCH ranges stay static)::

    GPU            title row 1, header row 2, data rows 3..62
    LLM models     title row 64, header row 65, data rows 66..265
    Quantization   title row 267, header row 268, data rows 269..288
"""

from __future__ import annotations

from copy import copy
from typing import Any, Iterable

from openpyxl.worksheet.worksheet import Worksheet

from core.methodology_constants import QUANTIZATION_FORMATS

GPU_ROWS = (3, 62)
MODEL_ROWS = (66, 265)
QUANT_ROWS = (269, 288)

GPU_TITLE_ROW, GPU_HEADER_ROW = 1, 2
MODEL_TITLE_ROW, MODEL_HEADER_ROW = 64, 65
QUANT_TITLE_ROW, QUANT_HEADER_ROW = 267, 268

GPU_HEADER = ["GPU", "Memory (GiB)", "FLOPS FP16 (TFLOPS)", "BW_GPU (GB/s)", "Memory type", "id"]
MODEL_HEADER = [
    "LLM",
    "P_total (B)",
    "P_active (B)",
    "L",
    "H",
    "N_attention",
    "N_kv",
    "head_dim",
    "Arch",
    "TS_max",
    "kv_lora_rank",
    "qk_rope_head_dim",
    "P_dense (B)",
    "P_moe (B)",
    "k_moe",
    "N_experts",
]
QUANT_HEADER = ["Формат", "B_quant (байт)", "Описание"]

COLUMNS = "ABCDEFGHIJKLMNOP"

# Datacenter-grade subset for the dropdown: complete numbers only.
MIN_GPU_MEMORY_GB = 24
MIN_GPU_TFLOPS = 100
MIN_GPU_BW_GBS = 600


def gpu_display_name(gpu: dict[str, Any]) -> str:
    """ "<vendor> <model_name>" without repeating the vendor ("AMD AMD Instinct ...")."""
    vendor = (gpu.get("vendor") or "").strip()
    name = (gpu.get("model_name") or "").strip()
    if vendor and name.lower().startswith(vendor.lower()):
        return name
    return f"{vendor} {name}".strip()


def datacenter_gpus(catalog: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for gpu in catalog:
        try:
            mem = float(gpu.get("memory_gb") or 0)
            tflops = float(gpu.get("tflops_fp16") or 0)
            bw = float(gpu.get("memory_bandwidth_gbs") or 0)
        except TypeError, ValueError:
            continue
        if mem >= MIN_GPU_MEMORY_GB and tflops >= MIN_GPU_TFLOPS and bw >= MIN_GPU_BW_GBS:
            out.append(gpu)
    out.sort(key=lambda g: (str(g.get("vendor") or ""), gpu_display_name(g)))
    return out


def model_arch_label(model: dict[str, Any]) -> str:
    moe = "MoE" if model.get("is_moe") else "dense"
    if model.get("is_mla") or (model.get("kv_lora_rank") or 0) > 0:
        attention = "MLA"
    else:
        n_attn = model.get("num_attention_heads") or 0
        n_kv = model.get("num_kv_heads") or n_attn
        attention = "MQA" if n_kv == 1 else ("GQA" if n_kv < n_attn else "MHA")
    return f"{moe}/{attention}"


def gpu_row(gpu: dict[str, Any]) -> list[Any]:
    return [
        gpu_display_name(gpu),
        float(gpu.get("memory_gb") or 0),
        float(gpu.get("tflops_fp16") or 0),
        float(gpu.get("memory_bandwidth_gbs") or 0),
        gpu.get("memory_type"),
        gpu.get("id"),
    ]


def model_row(model: dict[str, Any]) -> list[Any]:
    n_attn = model.get("num_attention_heads") or 0
    hidden = model.get("hidden_size") or 0
    head_dim = model.get("head_dim") or (hidden // n_attn if n_attn else 0)
    return [
        model.get("name"),
        model.get("params_total_b"),
        model.get("params_active_b") or model.get("params_total_b"),
        model.get("layers"),
        hidden,
        n_attn,
        model.get("num_kv_heads") or n_attn,
        head_dim,
        model_arch_label(model),
        model.get("max_context"),
        model.get("kv_lora_rank") or 0,
        model.get("qk_rope_head_dim") or 0,
        model.get("params_dense_b") or 0,
        model.get("params_moe_b") or 0,
        model.get("k_moe") or 0,
        model.get("n_experts") or 0,
    ]


def quant_row(q: dict[str, Any]) -> list[Any]:
    return [q["label"], q["bytes_per_param"], q["description"]]


def _write_block(
    ws: Worksheet,
    title_row: int,
    header_row: int,
    rows: tuple[int, int],
    title: str,
    header: list[str],
    data: list[list[Any]],
    title_style: Any,
    header_style: Any,
    data_style: Any,
) -> None:
    first, last = rows
    if len(data) > last - first + 1:
        raise ValueError(f"Reference block '{title}' overflow: {len(data)} > {last - first + 1}")
    ws.cell(title_row, 1).value = title
    if title_style is not None:
        ws.cell(title_row, 1)._style = copy(title_style)
    for col, value in enumerate(header, start=1):
        cell = ws.cell(header_row, col)
        cell.value = value
        if header_style is not None:
            cell._style = copy(header_style)
    for idx, row in enumerate(range(first, last + 1)):
        values = data[idx] if idx < len(data) else []
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row, col)
            cell.value = values[col - 1] if col - 1 < len(values) else None
            if data_style is not None:
                cell._style = copy(data_style)


def write_reference(
    ws: Worksheet,
    gpu_catalog: Iterable[dict[str, Any]],
    llm_catalog: Iterable[dict[str, Any]],
    styles: dict[str, Any] | None = None,
) -> None:
    """(Re)write the three catalog blocks of the Reference sheet."""
    styles = styles or {}
    title, header, data = styles.get("title"), styles.get("header"), styles.get("data")
    _write_block(
        ws,
        GPU_TITLE_ROW,
        GPU_HEADER_ROW,
        GPU_ROWS,
        "Каталог GPU (генерируется из каталога API: gpu_data.json + gpu_overrides.json)",
        GPU_HEADER,
        [gpu_row(g) for g in datacenter_gpus(gpu_catalog)],
        title,
        header,
        data,
    )
    _write_block(
        ws,
        MODEL_TITLE_ROW,
        MODEL_HEADER_ROW,
        MODEL_ROWS,
        "Каталог LLM моделей (генерируется из llm_data.json)",
        MODEL_HEADER,
        [model_row(m) for m in llm_catalog],
        title,
        header,
        data,
    )
    _write_block(
        ws,
        QUANT_TITLE_ROW,
        QUANT_HEADER_ROW,
        QUANT_ROWS,
        "Квантизация весов (методика §3.1, core/methodology_constants.py)",
        QUANT_HEADER,
        [quant_row(q) for q in QUANTIZATION_FORMATS],
        title,
        header,
        data,
    )


def capture_styles(ws: Worksheet) -> dict[str, Any]:
    """Styles of the current layout's title / header / data rows."""
    return {
        "title": copy(ws.cell(GPU_TITLE_ROW, 1)._style),
        "header": copy(ws.cell(GPU_HEADER_ROW, 1)._style),
        "data": copy(ws.cell(GPU_ROWS[0], 2)._style),
    }
